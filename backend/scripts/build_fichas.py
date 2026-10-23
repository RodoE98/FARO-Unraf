"""
build_fichas.py - Convierte el documento de fichas docentes (.docx) en JSON limpio.

Uso (desde la carpeta backend/):
    pip install python-docx            # solo hace falta para este script, no para el servidor
    python scripts/build_fichas.py
    python scripts/build_fichas.py ruta/al/Fichas.docx fichas/fichas_derivadas.json

Que hace:
  1. Lee el .docx (una tabla de cabecera con ID + titulo y una tabla de contenido por ficha).
  2. Extrae los campos de cada ficha (pregunta, respuesta, complemento, nota, ejemplo, ...).
  3. Convierte subindices y superindices a texto plano legible (f_x, e^x).
  4. VALIDA contra el propio documento: cantidad de fichas, orden de IDs contra el indice,
     campos obligatorios y etiquetas desconocidas. Si algo no cuadra, termina con error
     en lugar de generar un JSON a medias.

Cuando la catedra publique una version nueva del documento: reemplazar el .docx en fichas/,
volver a ejecutar este script y hacer commit del JSON resultante.
"""
import json
import os
import re
import sys

import docx
from docx.table import Table
from docx.text.paragraph import Paragraph

BASE_DIR = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
DEFAULT_DOCX = os.path.join(BASE_DIR, "fichas", "Fichas_completas_corpus_derivadas_FARO_UNRaf_v0.4.docx")
DEFAULT_JSON = os.path.join(BASE_DIR, "fichas", "fichas_derivadas.json")

# Etiqueta que aparece en negrita en el documento -> nombre del campo en el JSON
ETIQUETAS = {
    "Pregunta principal": "pregunta",
    "Respuesta curada": "respuesta",
    "Complemento conceptual": "complemento",
    "Nota de precisión": "nota",
    "Ejemplo": "ejemplo",
    "Prerrequisitos": "prerrequisitos",
    "Relacionadas": "relacionadas",
    "Origen en el corpus": "origen",
}
OBLIGATORIOS = ("pregunta", "respuesta", "prerrequisitos", "relacionadas", "origen")

RE_ID = re.compile(r"^[A-Z]{2,4}-\d{2}[A-Z]?$")
RE_INDICE = re.compile(r"^([A-Z]{2,4}-\d{2}[A-Z]?)\s+(.+)$")
RE_SECCION = re.compile(r"^(I|II|III|IV|V|VI|VII|VIII|IX|X)\.\s+(.+)$")


class ErrorFichas(Exception):
    pass


def _marcar(texto, tipo):
    """Subindice/superindice -> _x / ^x (con llaves si tiene mas de un caracter)."""
    t = texto.strip()
    if not t:
        return texto
    return f"{tipo}{t}" if len(t) == 1 else f"{tipo}{{{t}}}"


def texto_parrafo(p):
    """Texto plano de un parrafo, convirtiendo sub/superindices reales de Word.

    Word suele partir un mismo subindice en varios runs contiguos (p. ej. "(", "x,y", ")->(", "a,b", ")"
    en un limite). Se agrupan los runs consecutivos del mismo tipo antes de convertir:
    lim_{(x,y)->(a,b)} en lugar de un _x_ suelto por cada trozo.
    """
    grupos = []  # [[tipo, texto], ...]  tipo: "" (normal), "_" (subindice) o "^" (superindice)
    for r in p.runs:
        tipo = "_" if r.font.subscript else ("^" if r.font.superscript else "")
        if grupos and grupos[-1][0] == tipo:
            grupos[-1][1] += r.text
        else:
            grupos.append([tipo, r.text])
    partes = [t if tipo == "" else _marcar(t, tipo) for tipo, t in grupos]
    t = "".join(partes).replace("\u00a0", " ")
    return re.sub(r"[ \t]+", " ", t).strip()


def campos_de_celda(celda, id_ficha):
    """Devuelve {campo: texto} a partir de la celda de contenido de una ficha."""
    campos, actual = {}, None
    for p in celda.paragraphs:
        txt = texto_parrafo(p)
        if not txt:
            continue
        # Un parrafo que arranca con una etiqueta en negrita "Etiqueta:" abre un campo nuevo
        etiqueta = None
        if p.runs and p.runs[0].bold:
            m = re.match(r"^([^:]{2,40}):\s*(.*)$", txt, flags=re.S)
            if m:
                etiqueta, resto = m.group(1).strip(), m.group(2).strip()
        if etiqueta is not None:
            if etiqueta not in ETIQUETAS:
                raise ErrorFichas(f"{id_ficha}: etiqueta desconocida '{etiqueta}:' (¿cambió el formato del documento?)")
            actual = ETIQUETAS[etiqueta]
            if actual in campos:
                raise ErrorFichas(f"{id_ficha}: la etiqueta '{etiqueta}' aparece dos veces")
            campos[actual] = resto
        elif actual is not None:
            campos[actual] += "\n" + txt  # continuación del campo anterior
        else:
            raise ErrorFichas(f"{id_ficha}: texto antes de la primera etiqueta: '{txt[:60]}'")
    return campos


def parsear(ruta_docx):
    documento = docx.Document(ruta_docx)
    elementos = list(documento.element.body.iterchildren())

    titulo_doc, version = "", ""
    indice = []            # [(id, titulo)] en el orden del índice
    secciones = {}         # "I" -> "Prerrequisitos y funciones"
    fichas = []
    declaradas = {}        # sección -> cantidad declarada ("18 fichas temáticas.")
    apendice, en_apendice = [], False
    estado = "portada"     # portada -> indice -> cuerpo
    seccion_actual, ultimo_header = None, None

    for el in elementos:
        tag = el.tag.split("}")[1]
        if tag == "p":
            p = Paragraph(el, documento)
            txt = texto_parrafo(p)
            if not txt:
                continue
            estilo = p.style.name
            if estilo == "Corpus Title":
                titulo_doc = txt
            m_ver = re.search(r"versi[oó]n\s+([\d.]+)", txt, flags=re.I)
            if m_ver and estado == "portada":
                version = m_ver.group(1)
            if estilo == "Section Title" and txt.lower().startswith("índice"):
                estado = "indice"
                continue
            if estado == "indice":
                if estilo == "Section Title":
                    estado = "cuerpo"           # el primer título de sección real abre el cuerpo
                else:
                    m = RE_INDICE.match(txt)
                    if m:
                        indice.append((m.group(1), m.group(2).strip()))
                    continue
            if estado == "cuerpo":
                if estilo == "Section Title":
                    m = RE_SECCION.match(txt)
                    en_apendice = txt.lower().startswith("apéndice")
                    if m:
                        seccion_actual = m.group(1)
                        secciones[seccion_actual] = m.group(2).strip()
                    elif not en_apendice:
                        en_apendice = False
                    continue
                m_n = re.match(r"^(\d+)\s+fichas", txt)
                if m_n and seccion_actual:
                    declaradas[seccion_actual] = int(m_n.group(1))
                if en_apendice and txt.startswith("•"):
                    apendice.append(txt.lstrip("• ").strip())
        elif tag == "tbl" and estado == "cuerpo":
            tabla = Table(el, documento)
            if not tabla.rows:
                continue
            celdas = tabla.rows[0].cells
            if len(celdas) == 2 and RE_ID.match(celdas[0].text.strip()):
                titulo_h = " ".join(texto_parrafo(pp) for pp in celdas[1].paragraphs if texto_parrafo(pp))
                ultimo_header = (celdas[0].text.strip(), titulo_h)
            elif len(celdas) == 1 and celdas[0].text.strip().startswith("Pregunta principal"):
                if ultimo_header is None:
                    raise ErrorFichas("Tabla de contenido sin cabecera de ficha previa")
                id_f, titulo = ultimo_header
                ultimo_header = None
                campos = campos_de_celda(celdas[0], id_f)
                ficha = {
                    "id": id_f,
                    "titulo": titulo,
                    "seccion": seccion_actual,
                    "seccion_titulo": secciones.get(seccion_actual, ""),
                }
                for campo in ETIQUETAS.values():
                    ficha[campo] = campos.get(campo, "")
                fichas.append(ficha)

    # ---------------- VALIDACIONES ----------------
    errores = []
    ids = [f["id"] for f in fichas]
    if len(ids) != len(set(ids)):
        dup = sorted({i for i in ids if ids.count(i) > 1})
        errores.append(f"IDs duplicados: {dup}")
    if [i for i, _ in indice] != ids:
        solo_indice = [i for i, _ in indice if i not in ids]
        solo_cuerpo = [i for i in ids if i not in [x for x, _ in indice]]
        errores.append(f"El índice no coincide con las fichas. Solo en índice: {solo_indice}; solo en cuerpo: {solo_cuerpo}; "
                       f"o el orden difiere.")
    titulos_indice = dict(indice)
    for f in fichas:
        if titulos_indice.get(f["id"]) not in (None, f["titulo"]):
            errores.append(f"{f['id']}: título distinto en el índice ('{titulos_indice[f['id']]}') y en la ficha ('{f['titulo']}')")
        for campo in OBLIGATORIOS:
            if not f[campo].strip():
                errores.append(f"{f['id']}: falta el campo obligatorio '{campo}'")
    por_seccion = {}
    for f in fichas:
        por_seccion[f["seccion"]] = por_seccion.get(f["seccion"], 0) + 1
    for sec, n in declaradas.items():
        if por_seccion.get(sec, 0) != n:
            errores.append(f"Sección {sec}: el documento declara {n} fichas pero se leyeron {por_seccion.get(sec, 0)}")
    if not fichas:
        errores.append("No se leyó ninguna ficha")
    if errores:
        raise ErrorFichas("\n  - " + "\n  - ".join(errores))

    return {
        "titulo": titulo_doc,
        "version": version,
        "fuente": os.path.basename(ruta_docx),
        "secciones": secciones,
        "fichas": fichas,
        "temas_sin_cobertura": apendice,
    }


def main():
    ruta_docx = sys.argv[1] if len(sys.argv) > 1 else DEFAULT_DOCX
    ruta_json = sys.argv[2] if len(sys.argv) > 2 else DEFAULT_JSON
    try:
        datos = parsear(ruta_docx)
    except ErrorFichas as e:
        print(f"ERROR al leer {ruta_docx}:{e}")
        sys.exit(1)
    os.makedirs(os.path.dirname(ruta_json), exist_ok=True)
    with open(ruta_json, "w", encoding="utf-8") as f:
        json.dump(datos, f, ensure_ascii=False, indent=2)
        f.write("\n")
    print(f"OK: {len(datos['fichas'])} fichas (versión {datos['version']}) -> {ruta_json}")
    for sec, titulo in datos["secciones"].items():
        n = sum(1 for f in datos["fichas"] if f["seccion"] == sec)
        print(f"   {sec}. {titulo}: {n} fichas")
    print(f"   Temas sin cobertura (apéndice): {len(datos['temas_sin_cobertura'])}")


if __name__ == "__main__":
    main()
